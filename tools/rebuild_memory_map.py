#!/usr/bin/env python3
"""Replace the rows of 'Memory Map (Sec 8)' with the PDF extraction plus hand corrections.

Inputs : data/memory_map.json   (normalize_memory_map.py)
         data/corrections.csv   (location-keyed hand corrections; see corrections.py)
         the workbook itself     (previous version: Value Type and Reference Section are carried over
                                  for rows whose location is unchanged, and changes are reported)
Output : the workbook, in place, and ../Memory_Map_Rebuild_Changes.xlsx (what changed vs the
         previous version, plus the corrections log).
Then run update_param_reference.py for Key ID, Mask Bytes, SONiC, formatting and the Read Me.

Usage: python3 rebuild_memory_map.py ../CMIS_5.3_Parameter_Reference.xlsx [previous.xlsx]
"""
import collections, json, os, sys
import openpyxl
from copy import copy
import corrections

HERE = os.path.dirname(os.path.abspath(__file__))
MM = "Memory Map (Sec 8)"


def page_cell(p):
    return None if p == "Lwr" else "0x" + p


def page_key(v):
    return "Lwr" if v is None else v[2:]


def mask_str(masks):
    return " ".join("0x%02X" % m for m in masks)


def previous_rows(path):
    ws = openpyxl.load_workbook(path)[MM]
    h = [c.value for c in ws[1]]
    out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        d = dict(zip(h, r))
        n = d["Byte Length"] or 1
        d["_masks"] = [int(m, 16) for m in d["Mask Bytes"].split()] if d["Mask Bytes"] else [0xFF] * n
        out.append(d)
    return out


def main(path, prev_path):
    data = json.load(open(os.path.join(HERE, "data", "memory_map.json")))
    recs, corr_log = corrections.apply(data["records"], os.path.join(HERE, "data", "corrections.csv"))

    # merge identical records that the spec lists in two tables
    uniq = collections.OrderedDict()
    for r in recs:
        k = (r["page"], r["name"], r["byte"], tuple(r["masks"]), r["condition"])
        if k in uniq:
            if r["table"] not in uniq[k]["table"]:
                uniq[k]["table"] += "; " + r["table"]
        else:
            uniq[k] = dict(r)
    recs = list(uniq.values())

    prev = previous_rows(prev_path)
    pk = {}
    for d in prev:
        pk.setdefault((page_key(d["Page"]), d["CMIS Name"], d["Byte"], tuple(d["_masks"]), d.get("Condition")), d)
    matched, changes, rows = set(), [], []
    for r in recs:
        old = pk.get((r["page"], r["name"], r["byte"], tuple(r["masks"]), r["condition"]))
        vt, ref_sec = r["value_type"], None
        if old:
            matched.add(id(old))
            if old.get("Value Type") not in (None, "unspecified", "reserved"):
                vt = old["Value Type"]
            ref_sec = old.get("Reference Section")
        elif r["name"] != "Reserved":
            changes.append(("Added or moved", r["page"], r["name"],
                            f"byte {r['byte']} len {r['len']} mask {mask_str(r['masks'])}", r["table"], r["pdf_page"]))
        rows.append({"CMIS Name": r["name"], "Page": page_cell(r["page"]), "Byte": r["byte"],
                     "Byte Length": r["len"], "Mask Bytes": mask_str(r["masks"]), "Bank Scope": r["bank"],
                     "Value Type": None if r["name"] == "Reserved" else vt,
                     "Unit": r.get("unit"), "Scale": r.get("scale"), "Byte Order": r.get("byte_order"),
                     "Conversion Note": r.get("conv_note"), "Access": r["access"],
                     "Behavior": r["behavior"], "Requirement": r["requirement"], "Condition": r["condition"],
                     "Description": r["desc"], "Reference Section": ref_sec, "Reference Table": r["table"]})
    for d in prev:
        if id(d) not in matched and d["CMIS Name"] != "Reserved":
            changes.append(("Removed or moved", page_key(d["Page"]), d["CMIS Name"],
                            f"byte {d['Byte']} len {d['Byte Length']} mask {d['Mask Bytes']}", d["Reference Table"], None))

    def order(x):
        p = -1 if x["Page"] is None else int(x["Page"], 16)
        m = [int(v, 16) for v in x["Mask Bytes"].split()]
        return (p, x["Condition"] or "", x["Byte"], -max(m[0].bit_length(), 1), x["CMIS Name"] == "Reserved")
    rows.sort(key=order)

    wb = openpyxl.load_workbook(path)
    ws = wb[MM]
    hdr = [c.value for c in ws[1]]
    ws.delete_rows(2, ws.max_row)
    new_cols = [h for h in ("Unit", "Scale", "Byte Order", "Conversion Note") if h not in hdr]
    if new_cols:  # insert after Value Type, copying the header style
        at = hdr.index("Value Type") + 2
        ws.insert_cols(at, len(new_cols))
        for k, h in enumerate(new_cols):
            c = ws.cell(1, at + k, h)
            src = ws.cell(1, 1)
            c.font, c.fill, c.alignment, c.border = copy(src.font), copy(src.fill), copy(src.alignment), copy(src.border)
        # column widths do not shift with insert_cols; reassign by header name below (update script)
        hdr = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(hdr)}
    for i, r in enumerate(rows, start=2):
        for h, v in r.items():
            if h in col:
                ws.cell(i, col[h], v)
    try:
        wb.save(path)
    except PermissionError:
        raise SystemExit(f"Cannot write {path}: close it in Excel and re-run.")

    out = openpyxl.Workbook()
    s = out.active
    s.title = "Changes"
    s.append(["Change", "Page", "CMIS Name", "Detail", "Table", "PDF page"])
    for x in changes:
        s.append(list(x))
    c = out.create_sheet("Corrections log")
    c.append(["Status", "Page", "Byte", "Mask", "Expected name", "Action", "Field", "New value", "Detail", "Reason"])
    for st, k, detail in corr_log:
        c.append([st, k["page"], k["byte"], k["mask"], k["expected_name"], k["action"], k["field"], k["new_value"], detail, k["reason"]])
    for sh in (s, c):
        sh.auto_filter.ref = sh.dimensions
        sh.freeze_panes = "A2"
        for L, w in zip("ABCDEFGHIJ", (18, 6, 34, 40, 22, 9, 18, 12, 40, 60)):
            sh.column_dimensions[L].width = w
    rp = os.path.join(os.path.dirname(os.path.abspath(path)), "Memory_Map_Rebuild_Changes.xlsx")
    out.save(rp)
    print(f"{len(rows)} rows written; {len(changes)} changes vs previous; corrections: "
          f"{collections.Counter(x[0] for x in corr_log)} -> {rp}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2] if len(sys.argv) > 2 else sys.argv[1])
