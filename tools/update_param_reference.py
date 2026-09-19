#!/usr/bin/env python3
"""Post-process CMIS_5.3_Parameter_Reference.xlsx in place (idempotent).

Steps (Memory Map (Sec 8) unless noted):
  1. (Hand corrections are applied earlier, by rebuild_memory_map.py from data/corrections.csv.)
  2. Replace 'Stable ID' with 'Key ID' (rule below).
  3. Blank Mask Bytes when every byte is 0xFF (whole-byte / multi-byte fields).
  4. Add 'SONiC Cross-check' and 'SONiC Field' from data/sonic_cmis_fields.json
     (produce it with sonic_extract.py).
  5. Header-cell comments describing each column; filters on every kept sheet;
     Memory Map frozen at G2 (header row + columns A:F).

Key ID rule: Key ID = CMIS Name.  For names used more than once:
  a. if the duplicates span pages, append '.<page>' as 2-digit hex + 'h'
     (Lower Memory -> 'Lwr'), e.g. PageChecksum.01h;
  b. if still duplicated, append '.<byte>' in decimal;
  c. if still duplicated, append '.<mask>' as 0x-prefixed hex, e.g. Reserved.01h.158.0xF8.

Usage:  python3 update_param_reference.py ../CMIS_5.3_Parameter_Reference.xlsx
"""
import collections, csv, json, os, sys
import openpyxl
from openpyxl.comments import Comment
from copy import copy
from openpyxl.worksheet.table import TableColumn
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.formatting.rule import FormulaRule
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.styles import PatternFill

HERE = os.path.dirname(os.path.abspath(__file__))
MM = "Memory Map (Sec 8)"
LOWER = "Lwr"

MM_COMMENTS = {
    "Key ID": "Unique key per row. Equals CMIS Name; duplicates get suffixes in order: .<page> (2-digit hex + 'h', Lower Memory = 'Lwr') when duplicates span pages, then .<byte> (decimal), then .<mask> (0x hex). Not guaranteed stable across CMIS revisions.",
    "CMIS Name": "Field name as printed in OIF-CMIS-05.3 Section 8 tables (Reserved regions are named 'Reserved'). Hand corrections are listed in tools/data/corrections.csv.",
    "Page": "Upper Memory page as 0xNN. Blank = Lower Memory (bytes 0-127). 0x00 = Upper Page 00h.",
    "Byte": "First byte address (decimal, as in the spec). 0-127 Lower Memory, 128-255 Upper Memory.",
    "Byte Length": "Number of consecutive bytes occupied by the field.",
    "Mask Bytes": "Bit mask per occupied byte, increasing byte-address order, 0x hex. Blank = all bits of every occupied byte (whole-byte or multi-byte field).",
    "Bank Scope": "Blank (gray) = unbanked. Otherwise Lane-banked (value per bank of 8 lanes), Application-banked, or CDB-instance.",
    "Value Type": "Normalized storage/decoding type (bool, u8, u16, s16, bitfieldN, enumN, bitmaskN, ascii[N], F16, ...). Blank (gray) = Reserved region (CMIS Name 'Reserved'). 'unspecified' = not derivable from the table; the spec description is authoritative.",
    "Unit": "Engineering unit of the value after scaling (degC, uV, mV, uA, uW, dB, dBm, GHz, nm, ns, us, ms, m, km, W, %), as the spec states it. Blank = no physical unit (codes, flags, counters, ASCII) or unit depends on another register (see Conversion Note).",
    "Scale": "Value in Unit = raw value x Scale (raw interpreted per Value Type and Byte Order). E.g. TempMonValue: s16 x 0.00390625 = degC. Blank when not applicable or conditional.",
    "Byte Order": "For multi-byte numbers: Big-endian (CMIS default, most significant byte at the lowest address) or Little-endian where the spec says so (e.g. page 14h SNR and counters).",
    "Conversion Note": "Extra conversion rules: dependence on another register (Aux monitors, bias scaling factor, length multipliers), F16 encoding, or formulas.",
    "Access": "RO, RW, WO, RWW (as defined in the CMIS access-type table 8-3).",
    "Behavior": "Normal, ClearOnRead (latched flags), SelfClearing (write-triggered), or See reference.",
    "Requirement": "Rqd, Opt, Cnd, Adv when stated by the spec; blank when the table does not state one for the row.",
    "Condition": "Condition text for Cnd requirements (currently unpopulated).",
    "Description": "Field description text from the spec table (extraction; verify exact wording against the PDF for anything safety/timing critical).",
    "Reference Section": "Spec section referenced by the row, when present.",
    "Reference Table": "OIF-CMIS-05.3 table the row came from (e.g. 8-62). Titles: see the list of tables in the spec.",
    "SONiC Cross-check": "Location comparison with SONiC sonic-platform-common CmisMemMap/CCmisMemMap (bank 0, commit in tools/data/sonic_cmis_fields.json). Agrees = same bits; Agrees (SONiC splits) = SONiC covers exactly these bits with several fields; SONiC wider = a SONiC field contains these bits plus others; Partial overlap; Conflict: SONiC uses reserved bits = a SONiC field is defined on bits the spec marks reserved; Not in SONiC = page implemented but these bits are not; Page not in SONiC. Behavioral reference only, not a compliance oracle.",
    "SONiC Field": "Name(s) of the overlapping SONiC field(s); aliases at identical bits are joined with ' | '.",
}
TI_COMMENTS = {
    "Table #": "OIF-CMIS-05.3 table number (8-x Memory Map, 9-x CDB). Page-overview summary tables are excluded.",
    "Title": "Table title as printed in the spec.",
    "Section": "Memory Map (Section 8) or CDB (Section 9).",
}
CDB_COMMENTS = {
    "Table #": "OIF-CMIS-05.3 Section 9 table number.",
    "Table Title (Command)": "Table title, usually the CDB command.",
    "Page/Byte": "Leading address or token as printed (raw extraction).",
    "Field / Description (raw)": "Remaining row text, unsplit (raw extraction; not yet normalized).",
    "Type": "Type token when extracted.",
}


CDB = "CDB Commands"
CDB_COLS = [
    ("CMD ID", 10, "CDB command ID (hex), as written to 9Fh:128-129."),
    ("Command Title", 40, "Command title from the Section 9 overview table (or the detail table title when the command is missing from the overview)."),
    ("Group", 24, "Command group from Table 9-1 (ID ranges)."),
    ("Requirement", 12, "Rqd, Opt, Cnd or Adv as stated in the overview table. Blank = not stated (see Note). Whether a module supports an Adv/Opt command is advertised via the Capabilities Inquiry commands (0040h-0045h)."),
    ("Section", 10, "Spec section describing the command."),
    ("Detail Table", 12, "Spec table with the command's CMD/REPLY payload layout (fields are not in this workbook yet)."),
    ("Overview Table", 14, "Section 9 overview table that lists the command."),
    ("Description", 80, "Command description from the overview table (spec text)."),
    ("Note", 50, "Extraction note, e.g. a command defined by a detail table but not listed in any overview table."),
]


def write_cdb(wb, header_src):
    path = os.path.join(HERE, "data", "cdb_commands.json")
    if not os.path.exists(path):
        return None
    cmds = json.load(open(path))["commands"]
    if CDB in wb.sheetnames:
        del wb[CDB]
    s = wb.create_sheet(CDB)
    for j, (h, w, note) in enumerate(CDB_COLS, start=1):
        c = s.cell(1, j, h)
        c.font, c.fill, c.alignment = copy(header_src.font), copy(header_src.fill), copy(header_src.alignment)
        c.comment = comment(note)
        s.column_dimensions[openpyxl.utils.get_column_letter(j)].width = w
    keys = ["cmd_id", "title", "group", "requirement", "section", "detail_table", "overview_table", "description", "note"]
    for i, cmd in enumerate(cmds, start=2):
        for j, k in enumerate(keys, start=1):
            s.cell(i, j, cmd.get(k) or None)
    s.auto_filter.ref = s.dimensions
    s.freeze_panes = "B2"
    return s


def write_csv(sheet, csv_path):
    with open(csv_path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f, lineterminator="\n")
        for row in sheet.iter_rows(values_only=True):
            w.writerow(["" if v is None else v for v in row])


def scrub_office_metadata(wb):
    """Remove organisation-specific traces an Office install adds on save: sensitivity-label
    custom properties (MSIP_Label_*) and the label's page header/footer text; normalise the
    last-modified-by name to the document creator."""
    from openpyxl.worksheet.header_footer import HeaderFooter
    for s in wb.worksheets:
        s.HeaderFooter = HeaderFooter()
    try:
        props = wb.custom_doc_props
        for prop in list(props.props):
            if prop.name.startswith("MSIP_Label_"):
                props.props.remove(prop)
    except AttributeError:
        pass
    if wb.properties.creator:
        wb.properties.lastModifiedBy = wb.properties.creator


def comment(text):
    c = Comment(text, "CMIS tools")
    c.width, c.height = 320, max(80, 18 * (len(text) // 45 + 2))
    return c


def masks_of(mask, length):
    if not mask:
        return [0xFF] * (length or 1)
    return [int(m, 16) for m in str(mask).split()]


def fmt_page(p):
    return LOWER if p in (None, "") else str(p)[2:].upper() + "h"


def key_ids(recs):
    """recs: list of dicts with name, page, byte, mask(str). Returns list of keys."""
    keys = [r["name"] for r in recs]
    by = collections.defaultdict(list)
    for i, r in enumerate(recs):
        by[r["name"]].append(i)
    for name, idx in by.items():
        if len(idx) < 2:
            continue
        if len({recs[i]["page"] for i in idx}) > 1:
            for i in idx:
                keys[i] += "." + fmt_page(recs[i]["page"])
        for stage in ("byte", "mask"):
            cnt = collections.Counter(keys[i] for i in idx)
            for i in idx:
                if cnt[keys[i]] > 1:
                    if stage == "byte":
                        keys[i] += "." + str(recs[i]["byte"])
                    else:
                        keys[i] += "." + "_".join(
                            "0x%02X" % m for m in masks_of(recs[i]["mask"], recs[i]["len"]))
    dup = [k for k, c in collections.Counter(keys).items() if c > 1]
    if dup:
        raise SystemExit(f"Key ID collision(s) remain: {dup[:10]}")
    return keys


def sonic_status(rec, sonic_by_page, sonic_pages):
    page = None if rec["page"] in (None, "") else int(str(rec["page"]), 16)
    if page not in sonic_pages:
        return "Page not in SONiC", None
    bits = {(rec["byte"] + d, b) for d, m in enumerate(masks_of(rec["mask"], rec["len"]))
            for b in range(8) if m >> b & 1}
    hits = [f for f in sonic_by_page.get(page, []) if bits & f["cells"]]
    if not hits:
        return "Not in SONiC", None
    # collapse aliases (same bits, different names)
    groups = collections.OrderedDict()
    for f in hits:
        groups.setdefault(frozenset(f["cells"]), []).append(f["name"])
    names = "; ".join(" | ".join(sorted(set(v))) for v in groups.values())
    cellsets = list(groups)
    if rec["kind"] == "Reserved":
        if not any(c <= bits for c in cellsets):  # SONiC only reads the containing byte(s)
            return "SONiC wider", names
        return "Conflict: SONiC uses reserved bits", names
    if any(c == bits for c in cellsets):
        exact = [" | ".join(sorted(set(groups[c]))) for c in cellsets if c == bits]
        return "Agrees", exact[0]
    if all(c <= bits for c in cellsets) and frozenset().union(*cellsets) == bits:
        return "Agrees (SONiC splits)", names
    if any(c > bits for c in cellsets):
        return "SONiC wider", names
    return "Partial overlap", names


def main(path):
    wb = openpyxl.load_workbook(path)
    ws = wb[MM]
    hdr = [c.value for c in ws[1]]
    if hdr[0] == "Stable ID":
        ws.cell(1, 1).value = "Key ID"
        hdr[0] = "Key ID"
    col = {h: i + 1 for i, h in enumerate(hdr)}
    widths = {h: ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width for h, i in col.items()}

    # Reserved rows: blank Value Type ('reserved'); then drop Record Kind (CMIS Name 'Reserved' identifies them)
    for r in range(2, ws.max_row + 1):
        vt = ws.cell(r, col["Value Type"])
        if vt.value == "reserved" or ("Record Kind" in col and ws.cell(r, col["Record Kind"]).value == "Reserved"):
            vt.value = None
        bs = ws.cell(r, col["Bank Scope"])
        if bs.value == "Unbanked":
            bs.value = None
    for gone in ("Record Kind", "Validation Status"):  # retired columns
        if gone in col:
            ws.delete_cols(col[gone])
            col = {h: i + 1 for i, h in enumerate(c.value for c in ws[1])}
    for h, i in col.items():
        if widths.get(h):
            ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = widths[h]
    for h in ("SONiC Cross-check", "SONiC Field"):
        if h not in col:
            c = ws.cell(1, ws.max_column + 1, h)
            src = ws.cell(1, 1)
            c.font, c.fill, c.alignment, c.border = copy(src.font), copy(src.fill), copy(src.alignment), copy(src.border)
            col[h] = c.column

    rows = range(2, ws.max_row + 1)
    g = lambda r, h: ws.cell(r, col[h]).value
    recs = [dict(name=g(r, "CMIS Name"), page=g(r, "Page"), byte=g(r, "Byte"),
                 len=g(r, "Byte Length"), mask=g(r, "Mask Bytes"), kind="Reserved" if g(r, "CMIS Name") == "Reserved" else "Field") for r in rows]

    # 2. Key ID
    for r, k in zip(rows, key_ids(recs)):
        ws.cell(r, col["Key ID"]).value = k

    # 3. Mask Bytes
    blanked = 0
    for r, rec in zip(rows, recs):
        if rec["mask"] and all(m == 0xFF for m in masks_of(rec["mask"], rec["len"])):
            ws.cell(r, col["Mask Bytes"]).value = None
            blanked += 1

    # 4. SONiC
    sj = json.load(open(os.path.join(HERE, "data", "sonic_cmis_fields.json")))
    sonic_by_page = collections.defaultdict(list)
    for f in sj["fields"]:
        f["cells"] = frozenset(tuple(x) for x in f["bits"])
        sonic_by_page[f["page"]].append(f)
    sonic_pages = set(sonic_by_page)
    stat = collections.Counter()
    for r, rec in zip(rows, recs):
        s, n = sonic_status(rec, sonic_by_page, sonic_pages)
        ws.cell(r, col["SONiC Cross-check"]).value = s
        ws.cell(r, col["SONiC Field"]).value = n
        stat[s] += 1
    for c in ws[1]:
        if c.value not in MM_COMMENTS:
            c.comment = None
    ws.column_dimensions[openpyxl.utils.get_column_letter(col["SONiC Cross-check"])].width = 24
    ws.column_dimensions[openpyxl.utils.get_column_letter(col["SONiC Field"])].width = 36
    ws.column_dimensions["A"].width = 40

    # 5. comments, filters, freeze
    def decorate(sheet, comments, freeze):
        for c in sheet[1]:
            if c.value in comments:
                c.comment = comment(comments[c.value])
        if sheet.tables:
            # Excel forbids a sheet AutoFilter on top of a Table; the Table supplies
            # the filter buttons. Resize it and resync its column names to the header.
            sheet.auto_filter.ref = None
            for t in sheet.tables.values():
                t.ref = sheet.dimensions
                t.tableColumns = [TableColumn(id=i + 1, name=str(c.value)) for i, c in enumerate(sheet[1])]
                t.autoFilter = AutoFilter(ref=t.ref)
        else:
            sheet.auto_filter.ref = sheet.dimensions
        from openpyxl.worksheet.views import Selection
        sheet.sheet_view.pane = None          # avoid piling up duplicate pane selections on every run
        sheet.sheet_view.selection = [Selection()]
        sheet.freeze_panes = freeze
    WIDTHS = {"Key ID": 40, "CMIS Name": 30, "Page": 8, "Byte": 7, "Byte Length": 8, "Mask Bytes": 16,
              "Bank Scope": 17, "Value Type": 12, "Unit": 8, "Scale": 11, "Byte Order": 12,
              "Conversion Note": 45, "Access": 9, "Behavior": 13, "Requirement": 12, "Condition": 26,
              "Description": 72, "Reference Section": 11, "Reference Table": 11,
              "SONiC Cross-check": 24, "SONiC Field": 36}
    # reset column widths: stale grouped entries from earlier layouts (e.g. min=15 max=16) overlap
    # the new columns and make Excel mis-render them
    from openpyxl.worksheet.dimensions import DimensionHolder, ColumnDimension
    ws.column_dimensions = DimensionHolder(worksheet=ws)
    for h, i in col.items():
        L = openpyxl.utils.get_column_letter(i)
        ws.column_dimensions[L] = ColumnDimension(ws, index=L, width=WIDTHS.get(h, 14))
    decorate(ws, MM_COMMENTS, "G2")
    # Shading: Reserved rows 5% gray; 15% gray (Background 1, darker 15%) on meaningful blanks
    ws.conditional_formatting = ConditionalFormattingList()
    gray = PatternFill(fill_type="solid", start_color="FFD9D9D9", end_color="FFD9D9D9")
    light = PatternFill(fill_type="solid", start_color="FFF2F2F2", end_color="FFF2F2F2")
    last = ws.max_row
    # 1st rule: whole Reserved rows in 5% gray (stops here, so their blanks are not shaded darker)
    N = openpyxl.utils.get_column_letter(col["CMIS Name"])
    lastcol = openpyxl.utils.get_column_letter(ws.max_column)
    ws.conditional_formatting.add(f"A2:{lastcol}{last}",
                                  FormulaRule(formula=[f'${N}2="Reserved"'], fill=light, stopIfTrue=True))
    # then 15% gray on meaningful blanks in field rows
    for h in ("Page", "Mask Bytes", "Bank Scope", "Value Type", "Access", "Behavior", "Requirement"):
        L = openpyxl.utils.get_column_letter(col[h])
        ws.conditional_formatting.add(f"{L}2:{L}{last}", FormulaRule(formula=[f"LEN(${L}2)=0"], fill=gray))
    if "Table Index" in wb.sheetnames:
        decorate(wb["Table Index"], TI_COMMENTS, "A2")
    if "CDB Commands (Sec 9)" in wb.sheetnames:
        decorate(wb["CDB Commands (Sec 9)"], CDB_COMMENTS, "A2")

    cdb_sheet = write_cdb(wb, ws.cell(1, 1))
    ncdb = (cdb_sheet.max_row - 1) if cdb_sheet else 0

    # Read Me (kept in sync with the data)
    if "Read Me" in wb.sheetnames:
        import readme
        names = [ws.cell(r, col["CMIS Name"]).value for r in rows]
        errata = len(list(csv.DictReader(open(os.path.join(HERE, "data", "corrections.csv"), newline=""))))
        readme.write(wb["Read Me"], {
            "rows": len(names), "reserved": names.count("Reserved"), "fields": len(names) - names.count("Reserved"),
            "sonic_sha": sj["sha"], "errata": errata, "cdb": ncdb,
            "sonic": ", ".join(f"{k} {v}" for k, v in stat.most_common())})
    scrub_office_metadata(wb)
    try:
        wb.save(path)
    except PermissionError:
        raise SystemExit(f"Cannot write {path}: close it in Excel and re-run.")

    # plain-text copy of the Memory Map: diffable in git, easy for code and agents
    folder = os.path.dirname(os.path.abspath(path))
    write_csv(ws, os.path.join(folder, "CMIS_5.3_Memory_Map.csv"))
    if cdb_sheet:
        write_csv(cdb_sheet, os.path.join(folder, "CMIS_5.3_CDB_Commands.csv"))
    print(f"rows={len(recs)} masks_blanked={blanked} sonic={dict(stat)} sheets={wb.sheetnames}")


if __name__ == "__main__":
    main(sys.argv[1] if len(sys.argv) > 1 else os.path.join(HERE, "..", "CMIS_5.3_Parameter_Reference.xlsx"))
